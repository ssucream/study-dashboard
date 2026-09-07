"""업로드된 강의자료(PDF / PPTX / DOCX)에서 평문 텍스트를 추출한다.

이미지로 박힌 텍스트(스캔 PDF, 이미지 슬라이드)는 추출되지 않는다 — OCR은 스코프 밖.
추출 결과가 사실상 비어 있으면 호출부가 사용자에게 경고할 수 있도록 빈 문자열을 반환한다.
"""

from __future__ import annotations

import io
import re
from pathlib import Path

SUPPORTED_SUFFIXES = {".pdf", ".pptx", ".docx"}

# 업로드 1건 최대 크기 (20 MB).
MAX_MATERIAL_BYTES = 20 * 1024 * 1024


class MaterialExtractionError(RuntimeError):
    """지원하지 않는 형식이거나 파일이 손상된 경우."""


def is_supported_filename(filename: str) -> bool:
    return Path(filename).suffix.lower() in SUPPORTED_SUFFIXES


def _collapse_whitespace(text: str) -> str:
    text = text.replace("\x00", "")
    lines = [re.sub(r"[ \t\u00a0]+", " ", ln).strip() for ln in text.splitlines()]
    out: list[str] = []
    blank = 0
    for ln in lines:
        if ln:
            out.append(ln)
            blank = 0
        else:
            blank += 1
            if blank <= 1:
                out.append("")
    return "\n".join(out).strip()


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:  # pragma: no cover - 의존성 누락 환경
        raise MaterialExtractionError("pypdf 패키지가 설치되어 있지 않습니다.") from e

    try:
        reader = PdfReader(io.BytesIO(data))
        parts = [page.extract_text() or "" for page in reader.pages]
    except Exception as e:
        raise MaterialExtractionError(f"PDF를 읽지 못했습니다: {e}") from e
    return _collapse_whitespace("\n\n".join(parts))


def _extract_pptx(data: bytes) -> str:
    try:
        from pptx import Presentation
    except ImportError as e:  # pragma: no cover
        raise MaterialExtractionError("python-pptx 패키지가 설치되어 있지 않습니다.") from e

    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as e:
        raise MaterialExtractionError(f"PPTX를 읽지 못했습니다: {e}") from e

    parts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = "\n".join(p.text for p in shape.text_frame.paragraphs)
                if text.strip():
                    parts.append(text)
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
        notes = getattr(slide, "notes_slide", None) if slide.has_notes_slide else None
        if notes and notes.notes_text_frame and notes.notes_text_frame.text.strip():
            parts.append(notes.notes_text_frame.text)
        parts.append("")
    return _collapse_whitespace("\n".join(parts))


def _extract_docx(data: bytes) -> str:
    try:
        import docx
    except ImportError as e:  # pragma: no cover
        raise MaterialExtractionError("python-docx 패키지가 설치되어 있지 않습니다.") from e

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise MaterialExtractionError(f"DOCX를 읽지 못했습니다: {e}") from e

    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return _collapse_whitespace("\n".join(parts))


def extract_text(filename: str, data: bytes) -> str:
    """파일 확장자에 따라 텍스트를 추출한다. 지원하지 않으면 MaterialExtractionError."""
    if len(data) > MAX_MATERIAL_BYTES:
        raise MaterialExtractionError(
            f"파일이 너무 큽니다. 최대 {MAX_MATERIAL_BYTES // (1024 * 1024)}MB까지 지원합니다."
        )
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(data)
    if suffix == ".pptx":
        return _extract_pptx(data)
    if suffix == ".docx":
        return _extract_docx(data)
    raise MaterialExtractionError(
        f"지원하지 않는 형식입니다: {suffix or '(확장자 없음)'} — PDF / PPTX / DOCX만 가능합니다."
    )


def extract_text_from_path(path: Path) -> str:
    return extract_text(path.name, path.read_bytes())
