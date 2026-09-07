"""src/quiz/materials.py — 강의자료 텍스트 추출 테스트."""

import io

import pytest

from src.quiz.materials import (
    MAX_MATERIAL_BYTES,
    MaterialExtractionError,
    extract_text,
    is_supported_filename,
)


def test_is_supported_filename():
    assert is_supported_filename("a.pdf")
    assert is_supported_filename("b.PPTX")
    assert is_supported_filename("c.docx")
    assert not is_supported_filename("d.txt")
    assert not is_supported_filename("e")


def test_extract_docx_roundtrip():
    import docx

    document = docx.Document()
    document.add_paragraph("첫 번째 문단입니다.")
    document.add_paragraph("두 번째 문단입니다.")
    buf = io.BytesIO()
    document.save(buf)

    text = extract_text("lecture.docx", buf.getvalue())
    assert "첫 번째 문단입니다." in text
    assert "두 번째 문단입니다." in text


def test_extract_pptx_roundtrip():
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "슬라이드 본문 텍스트"
    buf = io.BytesIO()
    prs.save(buf)

    text = extract_text("deck.pptx", buf.getvalue())
    assert "슬라이드 본문 텍스트" in text


def test_extract_unsupported_suffix_raises():
    with pytest.raises(MaterialExtractionError, match="지원하지 않는"):
        extract_text("notes.txt", b"hello")


def test_extract_oversized_raises():
    with pytest.raises(MaterialExtractionError, match="너무 큽니다"):
        extract_text("big.pdf", b"x" * (MAX_MATERIAL_BYTES + 1))


def test_extract_corrupt_pdf_raises():
    with pytest.raises(MaterialExtractionError):
        extract_text("broken.pdf", b"%PDF-1.4 not really a pdf")
